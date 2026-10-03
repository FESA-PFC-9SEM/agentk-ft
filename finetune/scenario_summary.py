"""
Per-scenario summary table for one evaluated model, re-scored from the raw
outputs finetune/run_scenarios.py saved (no model needed): model, scenario,
no. errors, no. detected/undetected, no. corrected/uncorrected -- one sheet
per evaluation setting plus a Summary sheet.

Settings, from whichever result files exist in the run folder:
  - Sampled (scenarios_results.xlsx: temperature 0.7, N runs averaged) or
    Greedy (scenarios_results_greedy.xlsx: temperature 0, 1 run);
  - strict correction (the whole patch applied all-or-nothing) or per-finding
    correction (each finding's own ops applied on their own);
  - with --lenient, the same again ignoring schema errors: findings and patch
    are scored even when the response breaks the output contract (e.g. a
    zero-shot base model writing "hardc***" evidence or "containers[0]"
    paths), to separate "didn't find it" from "found it but broke the format".

Usage:
    python -m finetune.scenario_summary runs/<run>
    python -m finetune.scenario_summary runs/zero-shot_<model> --lenient
"""

from __future__ import annotations

import argparse
import collections
import json
from pathlib import Path

import yaml

from finetune import run_scenarios
from finetune.evaluate import strip_fences
from dataset.detect import detect_file
from dataset.schema import RULE_IDS
from finetune.metrics import paths_related, prf, rule_confusions
from finetune.run_scenarios import _finding_matches_instance, load_test_cases, score_run

HEADER = [
    "model",
    "scenario",
    "no. errors",
    "no. detected errors",
    "no. undetected errors",
    "no. corrected errors",
    "no. uncorrected errors",
]
RESULT_FILES = {
    "Sampled": ("scenarios_results.xlsx", "temperature 0.7, runs averaged"),
    "Greedy": ("scenarios_results_greedy.xlsx", "greedy decoding (deterministic), 1 run"),
}
CORRECTION_KEYS = {
    "strict": ("corrected", "patch applied all-or-nothing"),
    "per finding": ("corrected_per_finding", "each finding's patch ops applied on their own"),
}


def lenient_parse(text: str) -> tuple[dict | None, list[str]]:
    """Any JSON object in the response, schema errors ignored; the four
    top-level lists default to empty and malformed entries are dropped."""
    try:
        obj = json.loads(strip_fences(text))
    except (json.JSONDecodeError, TypeError):
        return None, ["invalid JSON"]
    if not isinstance(obj, dict):
        return None, ["not a JSON object"]
    for key in ("findings", "patch", "new_resources", "notes"):
        if not isinstance(obj.get(key), list):
            obj[key] = []
    obj["findings"] = [
        {**f, "message": str(f.get("message", "")), "evidence": str(f.get("evidence", ""))}
        for f in obj["findings"]
        if isinstance(f, dict) and "rule_id" in f
    ]
    obj["patch"] = [op for op in obj["patch"] if isinstance(op, dict) and "path" in op and "op" in op]
    return obj, []


def score_file(results_path: Path, cases: dict, lenient: bool) -> dict[str, dict]:
    """{scenario file: {"runs", "detected", "corrected", "corrected_per_finding"}} summed over runs."""
    from openpyxl import load_workbook

    rows = list(load_workbook(results_path, read_only=True)["Runs"].iter_rows(values_only=True))
    header = rows[0]
    strict_parser = run_scenarios.parse_model_output
    if lenient:
        run_scenarios.parse_model_output = lenient_parse
    try:
        totals: dict[str, dict] = collections.defaultdict(collections.Counter)
        for values in rows[1:]:
            row = dict(zip(header, values))
            name = row["file"]
            if name not in cases:
                continue
            docs = list(yaml.safe_load_all((Path("scenarios") / name).read_text(encoding="utf-8")))
            result = score_run(docs, cases[name], row["raw_output"] or "")
            t = totals[name]
            t["runs"] += 1
            for key in ("detected", "corrected", "corrected_per_finding"):
                t[key] += sum(bool(v.get(key)) for v in result["instances"].values())
        return totals
    finally:
        run_scenarios.parse_model_output = strict_parser


def classification_metrics(results_path: Path, cases: dict, lenient: bool) -> dict:
    """Precision/recall/F1 (finding level) and accuracy/precision/recall/F1
    (scenario x rule level) for one result file, pooled over all runs.

    Finding level: a model finding is correct if it matches a ground-truth
    error, or a real defect the ground truth doesn't list (confirmed by the
    project's detectors -- e.g. KSEC-003 on 8-newrelic's host agent); any
    other finding is a false positive. Recall is reported over the in-scope
    errors and over all errors (out-of-scope ones counted as misses).

    Scenario x rule: for each scenario run and each rule, "does this file
    have a KSEC-00X problem?" -- truth from the ground truth plus the
    detectors, prediction from the model's findings."""
    from openpyxl import load_workbook

    rows = list(load_workbook(results_path, read_only=True)["Runs"].iter_rows(values_only=True))
    header = rows[0]
    parse = lenient_parse if lenient else run_scenarios.parse_model_output
    correct = wrong = 0
    gt_hit_scope = gt_scope = gt_hit_all = gt_all = 0
    pairs = []
    for values in rows[1:]:
        row = dict(zip(header, values))
        name = row["file"]
        if name not in cases:
            continue
        docs = list(yaml.safe_load_all((Path("scenarios") / name).read_text(encoding="utf-8")))
        response, _ = parse(row["raw_output"] or "")
        findings = response.get("findings", []) if response else []
        real = [f.to_dict() for f in detect_file(docs)]
        instances = cases[name]
        for finding in findings:
            if any(i["rule_id"] and _finding_matches_instance(finding, i) for i in instances) or any(
                r["rule_id"] == finding.get("rule_id")
                and r["doc"] == finding.get("doc", 0)
                and paths_related(r["path"], str(finding.get("path", "")))
                for r in real
            ):
                correct += 1
            else:
                wrong += 1
        for inst in instances:
            hit = inst["rule_id"] is not None and any(_finding_matches_instance(f, inst) for f in findings)
            gt_all += 1
            gt_hit_all += hit
            if inst["rule_id"] is not None:
                gt_scope += 1
                gt_hit_scope += hit
        actual = {i["rule_id"] for i in instances if i["rule_id"]} | {r["rule_id"] for r in real}
        pairs.append((actual, {f.get("rule_id") for f in findings}))
    precision = round(correct / (correct + wrong), 4) if correct + wrong else None
    out = {"findings_reported": correct + wrong, "correct_findings": correct, "false_positives": wrong, "precision": precision}
    for label, hit, total in (("in scope", gt_hit_scope, gt_scope), ("all errors", gt_hit_all, gt_all)):
        recall = round(hit / total, 4) if total else None
        f1 = round(2 * precision * recall / (precision + recall), 4) if precision and recall else 0.0
        out[label] = {"errors": total, "detected": hit, "recall": recall, "f1": f1}
    out["rule_level"] = rule_confusions(pairs, sorted(RULE_IDS))
    return out


def build_workbook(run_dir: Path, lenient: bool, test_cases: Path = Path("scenarios/test_cases.yaml")):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    cases = load_test_cases(test_cases)
    model = run_dir.name
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["model", "setting", *HEADER[2:], "% detected", "% corrected"])

    def style(ws, widths):
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        for col, width in zip("ABCDEFGHI", widths):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"

    parsings = [False, True] if lenient else [False]
    lines = []
    for decoding, (file_name, decoding_desc) in RESULT_FILES.items():
        path = run_dir / file_name
        if not path.exists():
            continue
        for use_lenient in parsings:
            totals = score_file(path, cases, use_lenient)
            for correction, (key, correction_desc) in CORRECTION_KEYS.items():
                title = f"{decoding} - {correction}" + (" (lenient)" if use_lenient else "")
                ws = wb.create_sheet(title[:31])
                ws.append(HEADER)
                n_all = det_all = cor_all = 0.0
                for name in sorted(cases, key=lambda n: int(n.split("-")[0])):
                    n = len(cases[name])
                    t = totals.get(name, collections.Counter(runs=1))
                    runs = t["runs"] or 1
                    det, cor = round(t["detected"] / runs, 1), round(t[key] / runs, 1)
                    ws.append([model, name, n, det, round(n - det, 1), cor, round(n - cor, 1)])
                    n_all += n
                    det_all += det
                    cor_all += cor
                ws.append(
                    [model, "TOTAL", n_all, round(det_all, 1), round(n_all - det_all, 1), round(cor_all, 1), round(n_all - cor_all, 1)]
                )
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True)
                ws.append([])
                note = f"Setting: {decoding_desc}; {correction_desc}."
                if use_lenient:
                    note += " Lenient: schema errors ignored, findings/patch scored as given."
                ws.append([note + " No. errors includes the out-of-scope errors (always undetected/uncorrected)."])
                style(ws, (48, 22, 10, 12, 12, 12, 12))
                summary.append(
                    [
                        model,
                        title,
                        n_all,
                        round(det_all, 1),
                        round(n_all - det_all, 1),
                        round(cor_all, 1),
                        round(n_all - cor_all, 1),
                        round(100 * det_all / n_all, 1),
                        round(100 * cor_all / n_all, 1),
                    ]
                )
                lines.append(f"{title:34} detected {100 * det_all / n_all:5.1f}%   corrected {100 * cor_all / n_all:5.1f}%")
    style(summary, (48, 34, 10, 12, 12, 12, 12, 11, 11))

    metrics_ws = wb.create_sheet("Metrics")
    metrics_ws.append(
        ["model", "setting", "level", "TP", "FP", "FN", "TN", "precision", "recall", "F1", "accuracy"]
    )
    rules_ws = wb.create_sheet("Metrics per rule")
    rules_ws.append(["model", "setting", "rule", "TP", "FP", "FN", "TN", "precision", "recall", "F1", "accuracy"])
    for decoding, (file_name, _) in RESULT_FILES.items():
        path = run_dir / file_name
        if not path.exists():
            continue
        for use_lenient in parsings:
            setting = decoding + (" (lenient)" if use_lenient else "")
            m = classification_metrics(path, cases, use_lenient)
            for label in ("in scope", "all errors"):
                g = m[label]
                metrics_ws.append(
                    [model, setting, f"finding ({label})", g["detected"], m["false_positives"], g["errors"] - g["detected"],
                     None, m["precision"], g["recall"], g["f1"], None]
                )
            for label, c in (("scenario x rule (micro)", m["rule_level"]["micro"]),):
                metrics_ws.append([model, setting, label, c["tp"], c["fp"], c["fn"], c["tn"], c["precision"], c["recall"], c["f1"], c["accuracy"]])
            mac = m["rule_level"]["macro"]
            metrics_ws.append([model, setting, "scenario x rule (macro)", None, None, None, None, mac["precision"], mac["recall"], mac["f1"], mac["accuracy"]])
            for rule, c in m["rule_level"]["per_rule"].items():
                rules_ws.append([model, setting, rule, c["tp"], c["fp"], c["fn"], c["tn"], c["precision"], c["recall"], c["f1"], c["accuracy"]])
            lines.append(
                f"{setting:20} findings: P={m['precision']} R(in scope)={m['in scope']['recall']} "
                f"F1={m['in scope']['f1']} | R(all)={m['all errors']['recall']} F1={m['all errors']['f1']} | "
                f"rule micro acc={m['rule_level']['micro']['accuracy']} F1={m['rule_level']['micro']['f1']}"
            )
    metrics_ws.append([])
    metrics_ws.append(["finding level: a finding is correct if it matches a ground-truth error or a real defect confirmed by the detectors; no TN exists, so no accuracy."])
    metrics_ws.append(["scenario x rule: one yes/no question per scenario run and rule (does this file have a KSEC-00X problem?); micro = pooled counts, macro = mean over rules."])
    style(metrics_ws, (48, 16, 28, 7, 7, 7, 7, 10, 10, 10, 10))
    style(rules_ws, (48, 16, 10, 7, 7, 7, 7, 10, 10, 10, 10))
    return wb, lines


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("run_dir", help="folder holding scenarios_results.xlsx and/or scenarios_results_greedy.xlsx")
    parser.add_argument("--lenient", action="store_true", help="also score ignoring schema errors")
    parser.add_argument("--output", default=None, help="default: <run_dir>/scenario_summary.xlsx")
    args = parser.parse_args(argv)
    run_dir = Path(args.run_dir)
    wb, lines = build_workbook(run_dir, args.lenient)
    output = Path(args.output) if args.output else run_dir / "scenario_summary.xlsx"
    wb.save(output)
    print("\n".join(lines))
    print(f"\nWritten to {output}")


if __name__ == "__main__":
    main()
