"""
Runs a fine-tuned checkpoint once against every scenarios/*.yaml file and
scores it against scenarios/test_cases.yaml -- the hand-written ground truth,
one entry per individual error (rule, document, and enough to tell that
error apart from other findings of the same rule; see that file's header).

One analysis, the way the model is used: deterministic (greedy) decoding,
and each fix applied on its own -- an error counts as corrected when
applying the operations of the finding that reports it makes the project's
own detectors (dataset/detect.py) stop flagging it.

Writes one workbook (default: next to the adapter, scenarios_results.xlsx):
  - "Results": per scenario, the number of errors, detected/undetected and
    corrected/uncorrected, plus five metrics -- precision, recall, F1,
    accuracy and the share of errors corrected;
  - "Outputs": the model's raw answer per scenario (enough to re-score).

Usage:
    finetune/.venv/bin/python -m finetune.run_scenarios --adapter runs/<run>/lora_adapter
"""

from __future__ import annotations

import argparse
from pathlib import Path

import yaml

from dataset.detect import detect_file
from dataset.schema import RULE_IDS, SYSTEM_PROMPT
from finetune.evaluate import apply_multidoc_patch, parse_model_output, patch_ops_for_finding
from finetune.metrics import paths_related, rule_confusions

RESULTS_HEADER = [
    "model",
    "scenario",
    "no. errors",
    "no. detected errors",
    "no. undetected errors",
    "no. corrected errors",
    "no. uncorrected errors",
]


def load_test_cases(path: Path) -> dict[str, list[dict]]:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def load_scenario_files(scenarios_dir: Path, test_cases: dict) -> list[Path]:
    return sorted(
        (scenarios_dir / name for name in test_cases),
        key=lambda p: int(p.name.split("-")[0]) if p.name.split("-")[0].isdigit() else p.name,
    )


def _finding_matches_instance(finding: dict, instance: dict) -> bool:
    if finding.get("rule_id") != instance["rule_id"]:
        return False
    if instance.get("doc") is not None and finding.get("doc") != instance["doc"]:
        return False
    match_type = instance.get("match_type", "rule_only")
    if match_type == "rule_only":
        return True
    keyword = instance["match_keyword"].lower()
    if match_type == "message":
        return keyword in str(finding.get("message", "")).lower()
    if match_type == "evidence_prefix":
        return str(finding.get("evidence", "")).lower().startswith(keyword)
    raise ValueError(f"unknown match_type: {match_type!r}")


def _corrected(instance: dict, docs: list[dict], findings: list[dict], patch: list[dict]) -> bool:
    """Apply only the operations of the model's findings for this error (the
    demo's "apply this fix") and check the detectors no longer flag it."""
    ops = []
    for finding in findings:
        if _finding_matches_instance(finding, instance):
            ops.extend(op for op in patch_ops_for_finding(finding, patch) if op not in ops)
    if not ops:
        return False
    patched_docs, applied_ok = apply_multidoc_patch(docs, ops)
    if not applied_ok:
        return False
    return not any(_finding_matches_instance(f.to_dict(), instance) for f in detect_file(patched_docs))


def score_run(docs: list[dict], instances: list[dict], raw_output: str) -> dict:
    """Scores one answer -- no model involved. Per error: detected and
    corrected; per reported finding: right (a ground-truth error, or a real
    defect the ground truth doesn't list) or a false positive."""
    response, errors = parse_model_output(raw_output)
    result = {
        "schema_valid": response is not None,
        "schema_errors": "; ".join(errors),
        "instances": {inst["id"]: {"detected": False, "corrected": False} for inst in instances},
        "right": 0,
        "false_positives": 0,
        "predicted_rules": set(),
        "raw_output": raw_output,
    }
    if response is None:
        return result
    findings = response.get("findings", [])
    patch = response.get("patch", [])
    real = [f.to_dict() for f in detect_file(docs)]
    for finding in findings:
        is_right = any(_finding_matches_instance(finding, i) for i in instances) or any(
            r["rule_id"] == finding.get("rule_id")
            and r["doc"] == finding.get("doc", 0)
            and paths_related(r["path"], str(finding.get("path", "")))
            for r in real
        )
        result["right" if is_right else "false_positives"] += 1
    result["predicted_rules"] = {f.get("rule_id") for f in findings}
    for instance in instances:
        detected = any(_finding_matches_instance(f, instance) for f in findings)
        result["instances"][instance["id"]] = {
            "detected": detected,
            "corrected": detected and _corrected(instance, docs, findings, patch),
        }
    return result


def analyze(scored: dict[str, dict], test_cases: dict) -> dict:
    """Per-scenario counts and the five overall metrics, from score_run's
    results per scenario file."""
    per_file = {}
    errors = detected = corrected = right = false_positives = 0
    rule_pairs = []
    for name, result in scored.items():
        instances = test_cases[name]
        n_det = sum(v["detected"] for v in result["instances"].values())
        n_cor = sum(v["corrected"] for v in result["instances"].values())
        per_file[name] = {"errors": len(instances), "detected": n_det, "corrected": n_cor}
        errors += len(instances)
        detected += n_det
        corrected += n_cor
        right += result["right"]
        false_positives += result["false_positives"]
        docs = list(yaml.safe_load_all(Path("scenarios", name).read_text(encoding="utf-8")))
        actual = {i["rule_id"] for i in instances} | {f.rule_id for f in detect_file(docs)}
        rule_pairs.append((actual, result["predicted_rules"]))
    precision = right / (right + false_positives) if right + false_positives else 0.0
    recall = detected / errors if errors else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    accuracy = rule_confusions(rule_pairs, sorted(RULE_IDS))["micro"]["accuracy"]
    metrics = {
        "Precision": precision,
        "Recall": recall,
        "F1-score": f1,
        "Accuracy": accuracy,
        "Errors corrected": corrected / errors if errors else 0.0,
    }
    return {"per_file": per_file, "totals": {"errors": errors, "detected": detected, "corrected": corrected}, "metrics": metrics}


def write_excel(scored: dict[str, dict], test_cases: dict, model: str, output_path: Path) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    analysis = analyze(scored, test_cases)
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    def bold_header(row_cells):
        for cell in row_cells:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws.append(RESULTS_HEADER)
    bold_header(ws[1])
    for name, c in analysis["per_file"].items():
        ws.append([model, name, c["errors"], c["detected"], c["errors"] - c["detected"], c["corrected"], c["errors"] - c["corrected"]])
    t = analysis["totals"]
    ws.append([model, "TOTAL", t["errors"], t["detected"], t["errors"] - t["detected"], t["corrected"], t["errors"] - t["corrected"]])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    ws.append([])
    ws.append(["Metric", "Value"])
    bold_header(ws[ws.max_row][:2])
    for metric, value in analysis["metrics"].items():
        ws.append([metric, value])
        ws.cell(ws.max_row, 2).number_format = "0.0%"
    for col, width in zip("ABCDEFG", (48, 22, 10, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = width

    out = wb.create_sheet("Outputs")
    out.append(["scenario", "valid answer", "detected errors (#)", "corrected errors (#)", "schema errors", "raw output"])
    bold_header(out[1])
    for name, result in scored.items():
        out.append(
            [
                name,
                result["schema_valid"],
                ", ".join(str(i) for i, v in result["instances"].items() if v["detected"]),
                ", ".join(str(i) for i, v in result["instances"].items() if v["corrected"]),
                result["schema_errors"],
                result["raw_output"],
            ]
        )
    for col, width in zip("ABCDEF", (22, 12, 20, 20, 30, 80)):
        out.column_dimensions[col].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return analysis


def run(args: argparse.Namespace) -> None:
    from finetune.evaluate import generate_response_text, load_model

    test_cases = load_test_cases(Path(args.test_cases))
    model, tokenizer = load_model(args.adapter, args.max_seq_length)
    scored = {}
    for path in load_scenario_files(Path(args.scenarios_dir), test_cases):
        raw_text = path.read_text(encoding="utf-8")
        output_text = generate_response_text(model, tokenizer, SYSTEM_PROMPT, raw_text, args.max_new_tokens)
        scored[path.name] = score_run(list(yaml.safe_load_all(raw_text)), test_cases[path.name], output_text)
        n_detected = sum(v["detected"] for v in scored[path.name]["instances"].values())
        print(f"[{path.name}] detected {n_detected}/{len(test_cases[path.name])}", flush=True)

    adapter_dir = Path(args.adapter)
    output = Path(args.output) if args.output else (adapter_dir.parent if adapter_dir.exists() else Path(".")) / "scenarios_results.xlsx"
    run_name = adapter_dir.parent.name if adapter_dir.exists() else args.adapter
    analysis = write_excel(scored, test_cases, run_name, output)
    for metric, value in analysis["metrics"].items():
        print(f"{metric:18} {value:.1%}")
    print(f"\nWritten to {output}")


def parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--adapter", required=True, help="path to a saved LoRA adapter (or a base model name)")
    parser.add_argument("--scenarios-dir", default="scenarios")
    parser.add_argument("--test-cases", default="scenarios/test_cases.yaml")
    parser.add_argument("--output", default=None, help="default: scenarios_results.xlsx next to the adapter")
    parser.add_argument("--max-seq-length", type=int, default=4096)
    parser.add_argument("--max-new-tokens", type=int, default=1280)
    return parser.parse_args(argv)


def main(argv=None) -> None:
    run(parse_args(argv))


if __name__ == "__main__":
    main()
